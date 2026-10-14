"""Build the CW report (Word) from the generated diagrams and model."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from model import TABLES, COLLECTIONS

BUILD = Path(sys.argv[1] if len(sys.argv) > 1 else "../build")
OUTFILE = BUILD / "STUDENTID_DDD.docx"

doc = Document()
FONT = "Arial"

# ---------------------------------------------------------------- styles
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 18), (2, 14), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = FONT
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
    h.paragraph_format.space_after = Pt(6)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)
        r.font.size = Pt(10)


add_page_number(sec)


def H(text, lvl=1):
    return doc.add_heading(text, lvl)


def P(text="", bold_lead=None, italic=False, size=None, align=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def B(text, bold_lead=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hexcol)
    tcPr.append(s)


def T(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A68")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            r = para.add_run(str(v))
            r.font.size = Pt(size)
    if widths:
        t.autofit = False
        tblPr = t._tbl.tblPr
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 567))); tw.set(qn("w:type"), "dxa")
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        tblPr.append(tw)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 567)))
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def CODE(text, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), "F3F5F8")
    pPr.append(s)
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Courier New"
        r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        if i < len(lines) - 1:
            r.add_break()
    return p


def new_section(width_cm, height_cm, landscape=False, margin=2.2):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(width_cm), Cm(height_cm)
    s.left_margin = s.right_margin = Cm(margin)
    s.top_margin = s.bottom_margin = Cm(min(margin, 2.0))
    return s


def figure_page(png, note, paper="A3"):
    """Put a large diagram on its own page, choosing the orientation that shows it largest."""
    w, h = Image.open(BUILD / png).size
    pw, ph = (42.0, 29.7) if paper == "A3" else (29.7, 21.0)
    margin = 1.2
    best = None
    for land in (True, False):
        W, Hh = (pw, ph) if land else (ph, pw)
        uw, uh = W - 2 * margin, Hh - 2 * margin - 1.6
        scale = min(uw / w, uh / h)
        if not best or scale > best[0]:
            best = (scale, land, W, Hh)
    scale, land, W, Hh = best
    new_section(W, Hh, landscape=land, margin=margin)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    small = BUILD / ("q_" + png)
    Image.open(BUILD / png).convert("RGB").quantize(colors=128, method=Image.Quantize.MEDIANCUT).save(small, optimize=True)
    p.add_run().add_picture(str(small), width=Cm(w * scale))
    P(note, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def back_to_text():
    new_section(21, 29.7)


# ================================================================ TITLE PAGE
for _ in range(5):
    doc.add_paragraph()
P("Database Fundamentals (6CS546)", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
t = P("", align=WD_ALIGN_PARAGRAPH.CENTER)
r = t.add_run("Pandemic Resilience System (PRS)")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
P("Database Design and Development", size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Conceptual ER model, relational logical model (MS SQL Server) and document logical model (MongoDB)",
  italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    doc.add_paragraph()
for k, v in [("Student name: ", "[Your full name]"), ("Student ID: ", "[Your student ID]"),
             ("Module: ", "6CS546 Database Fundamentals"), ("Date: ", "October 2026")]:
    P(v, bold_lead=k, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ================================================================ CONTENTS
H("Contents")
for line in ["1  Introduction", "    1.1  Scenario summary", "    1.2  Design approach", "    1.3  Requirements",
             "2  Conceptual Design", "    2.1  Conceptual ER model (Figure 1, detail views 1a to 1c)",
             "    2.2  Assumptions and design decisions", "    2.3  Validating the model against the scenario",
             "3  Logical Design", "    3.1  Choosing a store for each part of the data",
             "    3.2  Relational model (Figure 2, detail views 2a to 2c)",
             "    3.3  Document model (Figure 3)", "    3.4  Integrating the two models",
             "4  Conclusion", "References", "Appendix A  Implementation files and test evidence"]:
    p = P(line)
    p.paragraph_format.space_after = Pt(2)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ================================================================ 1 INTRODUCTION
H("1  Introduction")
H("1.1  Scenario summary", 2)
P("The Pandemic Resilience System (PRS) is a government platform for managing critical activities during a "
  "pandemic. It serves three groups of users. Government officials set purchase limits and day-of-week purchase "
  "schedules for critical items, enrol and verify merchants, and validate vaccination records. Merchants, which "
  "operate one or more stores, check each customer's eligibility before a sale, record every purchase of a "
  "restricted item and keep their stock levels up to date. Members of the public enrol themselves and their "
  "children under 18, locate stores with stock, see whether they may buy today, and upload vaccination evidence "
  "either as a PDF or as a machine-readable FHIR report such as the NHS Immunization sample.")
P("Every individual receives a PRS-Id that maps to one or more national identifiers (passport, driving licence, "
  "BRP). National identifiers and dates of birth are encrypted, only specific roles may decrypt them, and every "
  "access to sensitive data is audited.")

H("1.2  Design approach", 2)
P("The design was produced in four steps:")
B("The scenario was broken into numbered requirements (R1 to R18, Table 1) so coverage could be checked.", "Requirements. ")
B("One technology-independent ER model (Figure 1) captures every entity, attribute and relationship, with "
  "sensitive attributes marked.", "Conceptual model. ")
B("Each part of the model was assigned to the store that suits it best (Table 2): structured, "
  "transactional data with strict integrity goes to MS SQL Server; semi-structured, unstructured, append-only "
  "or highly variable data goes to MongoDB. This is a polyglot persistence design.", "Classification. ")
B("Both logical models (Figures 2 and 3) were derived from the conceptual model using identical names for the "
  "same things, and linked through shared identifiers (PRSId, UserId, StoreId, OfficialId, UploadId). Both were "
  "then implemented and tested (Appendix A) to prove the models work, including importing the NHS FHIR sample.",
  "Derivation. ")
P("All diagrams, the SQL Server DDL and the report tables are generated from one model definition, so a name "
  "cannot differ between the ERD, the schema and the collections.")

H("1.3  Requirements", 2)
REQS = [
    ("R1", "Assign every individual a unique PRS-Id", "Person"),
    ("R2", "Map the PRS-Id to one or more national identifiers of different types", "NationalIdentifier, IdentifierType"),
    ("R3", "Encrypt national identifiers and DOB; only authorised roles decrypt", "[E] attributes; Permission (DECRYPT_*)"),
    ("R4", "Enrol family members (children under 18), each with their own PRS-Id", "FamilyLink"),
    ("R5", "Purchase limits per item and period (e.g. 1 milk/day, 7 masks/week)", "PurchaseLimit"),
    ("R6", "Purchase schedules by DOB (birth-year digit to day of week)", "PurchaseSchedule, Person.BirthYearDigit"),
    ("R7", "Enrol merchants: licence number, name, operational addresses", "Merchant, Store"),
    ("R8", "Separate account for each operational location", "Store, UserAccount"),
    ("R9", "Record every purchase of critical items", "Purchase, PurchaseLine"),
    ("R10", "Maintain stock levels; locate the nearest store with stock", "StockLevel, Store.GeoLocation"),
    ("R11", "Check eligibility from DOB and purchase history before a sale", "PurchaseSchedule, PurchaseLimit, Purchase"),
    ("R12", "Alert officials when individuals or businesses break limits", "Alert"),
    ("R13", "Track vaccinations of residents and visitors; compliance per dose and period", "VaccinationRecord, Vaccine, VaccinationPolicy"),
    ("R14", "Upload records as PDF or FHIR JSON/XML; save machine-readable reports automatically", "VaccinationUpload, GridFS bucket PRSFiles"),
    ("R15", "Officials validate and verify records; use them at border checks", "Verification, BorderCheck"),
    ("R16", "Role-based access control", "UserAccount, Role, Permission"),
    ("R17", "Audit access to sensitive data", "AuditLog"),
    ("R18", "Notify people of nearby stock; publish real-time news and vaccination schedules", "Notification, Announcement"),
]
T(["Id", "Requirement (from the scenario)", "Entities that satisfy it"], REQS, widths=[1.2, 9.4, 6.0])
P("Table 1. Requirements and the entities that cover them.", italic=True, size=10)

# ================================================================ 2 CONCEPTUAL
H("2  Conceptual Design")
H("2.1  Conceptual ER model", 2)
P("Figure 1 is the complete conceptual ER model in Crow's Foot notation. It is technology-independent: it shows "
  "entities, attributes, identifiers and relationships with cardinality and participation, but no data types, "
  "tables or foreign keys. Entities are coloured by functional area. Attributes marked [E] must be stored "
  "encrypted and attributes marked [S] are sensitive personal data. Because the full model has 26 entities, "
  "Figures 1a to 1c repeat it as three detail views at a readable size; entities belonging to another view "
  "appear there as dashed grey boxes. Full-resolution SVG files of every figure are supplied with this report.")
P("Associative entities are used where a many-to-many relationship carries its own data: FamilyLink "
  "(guardian to dependant, with the relationship type), StockLevel (store to item, with quantity) and "
  "PurchaseLine (purchase to item, with quantity). The two M:N relationships that carry no data, UserAccount "
  "to Role and Role to Permission, are drawn as M:N and resolved in the logical model.")

figure_page("fig1_conceptual_erd.png", "Figure 1. PRS conceptual ER model (complete). Full-resolution file: fig1_conceptual_erd.svg")
figure_page("fig1a_conceptual_view.png", "Figure 1a. Detail view: people, identity, access and public information.")
figure_page("fig1b_conceptual_view.png", "Figure 1b. Detail view: merchants and supply chain.")
figure_page("fig1c_conceptual_view.png", "Figure 1c. Detail view: vaccination and border control.")
back_to_text()

H("2.2  Assumptions and design decisions", 2)
B("A person can hold several identifiers (for example a passport and a BRP), so Person to NationalIdentifier is "
  "1 to 1..N. IdentifierType is a separate entity so new identifier types can be added without changing the "
  "structure. One identifier can be flagged as primary.", "Identity. ")
B("A child can have two guardians and a guardian can enrol several children, so FamilyLink resolves an M:N "
  "recursive relationship on Person. Children may have no UserAccount of their own; the guardian acts for them.",
  "Family members. ")
B("DateOfBirth is encrypted, but schedules depend only on the last digit of the birth year. BirthYearDigit is "
  "derived once at enrolment and stored in clear, so a store can check the schedule without decrypting the DOB. "
  "It reveals one digit only and is a deliberate, documented redundancy. The under-18 rule for dependants is "
  "checked by the enrolment service, which is allowed to decrypt the DOB.", "Schedules without decrypting DOB. ")
B("A purchase limit applies to a person across all stores, so eligibility uses the person's whole purchase "
  "history. A week runs Monday to Sunday.", "Limits. ")
B("A business (Merchant) operates one or more locations (Store). Each store has one or more UserAccounts. A "
  "UserAccount belongs to exactly one of a Person, a Store or a GovernmentOfficial (an exclusive arc).",
  "Merchants and accounts. ")
B("An Alert concerns a person and/or a store, may be triggered by a purchase, and references the limit or "
  "schedule that was broken, so officials can see exactly which rule failed.", "Alerts. ")
B("VaccinationUpload is the evidence a person submits (FHIR or PDF). A machine-readable upload yields one "
  "VaccinationRecord per dose; a PDF yields none until an official reads it and keys the doses in. Each record "
  "is assessed against the VaccinationPolicy for its vaccine and dose, and every check by an official is kept as "
  "a Verification, so the history of who verified what is preserved.", "Vaccination evidence. ")
B("Vaccine is identified by its SNOMED CT product code, the same code FHIR uses in vaccineCode, so imported "
  "reports link to it without translation.", "Vaccine codes. ")
B("AuditLog records who did what to which record and when; Notification and Announcement support the public "
  "dashboard's real-time information.", "Audit and information. ")

H("2.3  Validating the model against the scenario", 2)
P("Each story below was traced through Figure 1 and then run against the implemented databases (Appendix A).")
B("The sample traveller was born in 1965 (BirthYearDigit 5). PurchaseSchedule allows digits 5 and 7 to buy face "
  "masks on Thursday, so a Monday attempt is refused and an Alert of type ScheduleBreach is raised. On Thursday "
  "five masks are sold (Purchase and PurchaseLine, StockLevel reduced). A further three that week would exceed "
  "PurchaseLimit (7 per week, summed from Purchase history), so they are refused with a LimitBreach alert.",
  "A pharmacy sells face masks. ")
B("The parent's Person is linked to the child's Person through FamilyLink. The parent's UserAccount submits a "
  "VaccinationUpload whose PRSId is the child's; the PDF is stored as a file. An official reviews it, creates "
  "VaccinationRecords with SourceType PDF and records a Verification.", "A parent uploads a child's PDF certificate. ")
B("The Patient's NHS number is matched to NationalIdentifier (type NHS_NUMBER) to find the PRS-Id and the birth "
  "date is compared with Person.DateOfBirth. The two Immunization resources become two VaccinationRecords "
  "linked to the upload, each assessed against VaccinationPolicy for its dose.", "A traveller uploads the NHS FHIR report. ")
B("The officer finds the person through the NationalIdentifier presented, checks the compliance of their "
  "verified VaccinationRecords and records a BorderCheck with its outcome.", "A border check. ")
B("StockLevel joined to Store (with its location) gives the nearest stores with stock; PurchaseSchedule and "
  "PurchaseLimit tell the person whether they may buy today; a Notification can tell them proactively.",
  "Finding a mask nearby. ")

# ================================================================ 3 LOGICAL
H("3  Logical Design")
H("3.1  Choosing a store for each part of the data", 2)
CLASS = [
    ("Person, NationalIdentifier, IdentifierType, FamilyLink", "SQL Server",
     "Core identity: must be unique, consistent and joined to almost everything; needs column encryption."),
    ("UserAccount, Role, Permission (+ UserRole, RolePermission)", "SQL Server",
     "Security rules need strict referential integrity; M:N is natural as junction tables."),
    ("Merchant, Store, CriticalItem, StockLevel", "SQL Server",
     "Structured master data; stock updates must be transactional; spatial index finds the nearest store."),
    ("PurchaseLimit, PurchaseSchedule, Purchase, PurchaseLine, Alert", "SQL Server",
     "Eligibility check, sale and stock change must happen in one ACID transaction; aggregate queries over history."),
    ("Vaccine, VaccinationPolicy, VaccinationRecord, Verification, BorderCheck", "SQL Server",
     "Normalised dose facts are queried at border checks and for compliance; must reference Person reliably."),
    ("VaccinationUpload (+ GridFS bucket PRSFiles)", "MongoDB",
     "FHIR bundles are nested and differ between countries; PDFs and images are unstructured binary files."),
    ("AuditLog", "MongoDB",
     "Very high write volume, append-only, details vary by action; time-series collection scales horizontally."),
    ("Notification, Announcement", "MongoDB",
     "Polymorphic payloads and attachments; short-lived, read-heavy feeds for the public dashboard."),
]
T(["Entities", "Store", "Reason"], CLASS, widths=[6.0, 2.4, 8.2])
P("Table 2. Data classification: relational versus document.", italic=True, size=10)
P("Structured data with strict integrity needs, transactions and many joins belongs in the relational model. "
  "Semi-structured data (FHIR), unstructured files (PDF and images) and high-volume, variable event data "
  "(audit, notifications) belong in the document model. Every conceptual entity appears in exactly one logical "
  "model, so no data is stored twice. The one exception is ExtractedDoses inside VaccinationUpload, which is a "
  "read-only snapshot of what was parsed from the source, kept as evidence.")

H("3.2  Relational logical model (MS SQL Server)", 2)
P("Figure 2 shows all 24 tables with their columns, SQL Server data types, primary keys (PK), foreign keys (FK), "
  "unique constraints (UQ), NOT NULL (NN) and Always Encrypted columns. Figures 2a to 2c are readable detail "
  "views. The complete DDL, including CHECK constraints, indexes and stored procedures, is in "
  "prs_sqlserver_schema.sql (Appendix A).")
figure_page("fig2_relational_schema.png", "Figure 2. PRS relational logical model (complete). Full-resolution file: fig2_relational_schema.svg")
figure_page("fig2a_relational_view.png", "Figure 2a. Detail view: people, identity and access.")
figure_page("fig2b_relational_view.png", "Figure 2b. Detail view: merchants and supply chain.")
figure_page("fig2c_relational_view.png", "Figure 2c. Detail view: vaccination and border control.")
back_to_text()

H("Mapping from the conceptual model", 3)
B("Each entity became a table with the same name, and its identifier became the primary key. PRSId keeps its "
  "business format (PRS + 9 digits, generated from a SEQUENCE). Other entities use IDENTITY surrogate keys, "
  "with natural keys such as BusinessLicenceNumber, StaffNumber and PermissionCode protected by UNIQUE "
  "constraints. Vaccine keeps its natural SNOMED code as the key because FHIR uses it directly.", "Entities. ")
B("Each 1:N relationship became a foreign key on the N side. Optional participation became a nullable FK "
  "(for example Alert.LimitId), mandatory participation a NOT NULL FK.", "1:N relationships. ")
B("UserAccount to Role and Role to Permission became the junction tables UserRole and RolePermission. The "
  "associative entities became tables with composite primary keys: FamilyLink (GuardianPRSId, DependentPRSId), "
  "StockLevel (StoreId, ItemId) and PurchaseLine (PurchaseId, ItemId).", "M:N relationships. ")
B("The exclusive arc on UserAccount is enforced by a CHECK constraint: a Public account must have PRSId only, a "
  "Merchant account StoreId only, an Official account OfficialId only.", "Exclusive arc. ")

H("Normalisation", 3)
P("All tables are in third normal form. There are no repeating groups: a purchase's items are rows in "
  "PurchaseLine, a schedule has one row per birth-year digit rather than a list such as '0,2', and a person's "
  "identifiers are rows in NationalIdentifier. There are no partial dependencies: in StockLevel and PurchaseLine "
  "the quantity depends on the whole composite key. There are no transitive dependencies: vaccine name and "
  "manufacturer depend on VaccineCode, so they live in Vaccine, not in VaccinationRecord, and identifier type "
  "names live in IdentifierType. Two controlled redundancies are kept on purpose: Person.BirthYearDigit "
  "(section 2.2) and VaccinationRecord.VerificationStatus, the current outcome of the latest Verification, "
  "which keeps border checks to a single-table lookup.")

H("Integrity and business rules", 3)
B("CHECK constraints restrict codes and ranges (PeriodType in Day/Week, BirthYearDigit 0 to 9, "
  "AllowedDayOfWeek 1 to 7, QuantityAvailable >= 0, EffectiveTo >= EffectiveFrom, GuardianPRSId <> DependentPRSId).",
  "Domain rules. ")
B("UNIQUE (IdentifierTypeCode, IssuingCountryCode, IdentifierValue) stops the same passport being mapped to two "
  "people. A filtered unique index on (PRSId, SourceResourceId) makes FHIR imports idempotent: uploading the same "
  "report twice cannot create duplicate doses.", "Uniqueness. ")
B("usp_CheckPurchaseEligibility evaluates the schedule and the limit against purchase history; usp_RecordPurchase "
  "runs the check, the sale and the stock decrement in one transaction and raises an Alert when a rule is "
  "broken; usp_FindNearestStock uses the spatial index on Store.GeoLocation; vw_VaccinationCompliance reports "
  "compliance per dose.", "Procedures. ")

H("Security in the relational model", 3)
B("NationalIdentifier.IdentifierValue uses deterministic Always Encrypted, so equality lookups still work "
  "(border checks, FHIR identity matching, the uniqueness constraint) while the server never sees plaintext. "
  "Person.DateOfBirth uses randomised encryption, which is stronger and never needs to be searched. Keys are held "
  "outside the database (e.g. Azure Key Vault) and released only to services acting for a role with "
  "DECRYPT_NATIONAL_ID or DECRYPT_DOB.", "Encryption ([E], AE-Det/AE-Rnd). ")
B("Other personal data (names, contact details, postcodes) is protected by Transparent Data Encryption at rest, "
  "Dynamic Data Masking for roles that do not need it, and TLS in transit.", "Sensitive data ([S]). ")
B("Roles and permissions are data (Role, Permission, RolePermission), so officials can change who may do what "
  "without code changes. A row-level security predicate based on FamilyLink lets a Citizen see only their own "
  "and their dependants' rows.", "RBAC. ")

H("3.3  Document logical model (MongoDB)", 2)
P("Figure 3 shows the four collections and the GridFS bucket, their fields and types, which parts are embedded "
  "(italic) and how each collection references SQL Server through shared keys (dashed). The validators, indexes "
  "and sample data are in prs_mongodb_setup.js.")
figure_page("fig3_document_model.png", "Figure 3. PRS document logical model. Full-resolution file: fig3_document_model.svg")
back_to_text()

H("Collections", 3)
B("One document per submission. A FHIR upload keeps the whole Bundle in OriginalContent exactly as received, "
  "whatever its country or version, and adds ExtractedDoses, the normalised fields PRS needs, each with a back "
  "reference to the VaccinationRecordId it created. A PDF or image upload embeds only file metadata; the binary "
  "is in GridFS. A $jsonSchema validator requires OriginalContent for FHIR formats and File for PDF/IMAGE.",
  "VaccinationUpload. ")
B("A time-series collection (timeField Timestamp, metaField Actor). The actor's roles are embedded as a snapshot "
  "so the log shows what rights the user had at that moment, even if their roles change later. The application "
  "connects with a MongoDB role that can only insert and read, so entries cannot be edited or deleted.", "AuditLog. ")
B("Each type carries a different Payload (stock nearby, limit warning, vaccination due, news). A TTL index "
  "deletes notifications after 90 days.", "Notification. ")
B("Public news and vaccination schedules with embedded attachment references and a text index for search.",
  "Announcement. ")
B("PDF certificates and images are stored as 255 KB chunks in PRSFiles.chunks with metadata (PRSId, purpose) in "
  "PRSFiles.files. Large files never bloat the upload documents, and a SHA-256 hash prevents the same file being "
  "registered twice.", "Unstructured files (GridFS bucket PRSFiles). ")

H("Embedding versus referencing", 3)
EMB = [
    ("OriginalContent (FHIR Bundle)", "Embed", "Owned by the upload, read together, small (well under 16 MB)."),
    ("ExtractedDoses[]", "Embed", "Bounded list owned by the upload; shows exactly what was parsed."),
    ("File metadata", "Embed", "One per upload, always shown with it."),
    ("File content (PDF/image)", "Reference (GridFS)", "Binary, can be large, streamed on demand."),
    ("Person", "Reference (PRSId)", "Owned by SQL Server, shared by everything, updated independently."),
    ("VaccinationRecord", "Reference (VaccinationRecordId)", "Two-way link between the evidence and the normalised dose."),
    ("Actor in AuditLog", "Embed snapshot", "Must be historically accurate, not the user's current roles."),
    ("PublishedBy in Announcement", "Embed snapshot", "Display name read with every announcement."),
    ("Payload in Notification", "Embed", "Different shape per type; never queried on its own."),
]
T(["Data", "Decision", "Reason"], EMB, widths=[5.0, 4.2, 7.4])
P("Table 3. Embedding and referencing decisions.", italic=True, size=10)

H("Example documents", 3)
P("The documents below were produced by importing the NHS sample (fhir_ingest.py) and a sample PDF "
  "(pdf_upload.py) into the implemented database. OriginalContent is shortened here; it holds the full bundle.")
ex = json.loads((BUILD / "example_documents.json").read_text())
fhir = ex["VaccinationUpload_FHIR"]
fhir["OriginalContent"] = {"resourceType": "Bundle", "type": "searchset", "total": 2,
                           "entry": ["... 2 Immunization + 1 Patient resources, exactly as received ..."]}
for d in fhir["ExtractedDoses"]:
    d["VaccineName"] = d["VaccineName"][:40] + "..."
    d.pop("Site", None)
    d.pop("Route", None)
fhir["ExtractedDoses"] = fhir["ExtractedDoses"]


def jdump(o):
    return json.dumps(o, indent=2, ensure_ascii=False)


P("VaccinationUpload created from the NHS FHIR report:", bold_lead="")
CODE(jdump(fhir))
P("VaccinationUpload for a PDF certificate (unstructured):")
pdf = ex["VaccinationUpload_PDF"]
CODE(jdump(pdf))
P("AuditLog entry written by the import:")
CODE(jdump(ex["AuditLog"]))
P("Notification telling a person about nearby stock:")
CODE(jdump(ex["Notification"]))

H("Schema definition (extract)", 3)
P("MongoDB is schema-flexible, but PRS enforces the fields it depends on with a $jsonSchema validator. The extract "
  "below shows the rules for VaccinationUpload; the full set is in prs_mongodb_setup.js.")
CODE("""db.createCollection("VaccinationUpload", { validator: { $jsonSchema: {
  bsonType: "object",
  required: ["PRSId", "UploadedByUserId", "Format", "UploadedAt", "ProcessingStatus"],
  properties: {
    PRSId:            { bsonType: "string", pattern: "^PRS[0-9]{9}$" },
    UploadedByUserId: { bsonType: "int" },
    Format:           { enum: ["FHIR_JSON", "FHIR_XML", "PDF", "IMAGE"] },
    ProcessingStatus: { enum: ["Received", "Parsed", "NeedsReview", "Failed"] },
    OriginalContent:  { bsonType: ["object", "binData"] },   // binData when CSFLE-encrypted
    File: { bsonType: "object",
            required: ["GridFsId", "FileName", "ContentType", "SizeBytes", "Sha256"] },
    ExtractedDoses: { bsonType: "array", items: { bsonType: "object",
            required: ["VaccineCode", "DoseNumber", "OccurrenceDateTime"] } }
  },
  oneOf: [ { properties: { Format: { enum: ["FHIR_JSON", "FHIR_XML"] } }, required: ["OriginalContent"] },
           { properties: { Format: { enum: ["PDF", "IMAGE"] } },          required: ["File"] } ]
}}, validationAction: "error" });""")

H("Handling messy imported data", 3)
P("The NHS sample shows why the original document must be kept and verified rather than trusted. Both "
  "Immunization resources say doseNumberPositiveInt = 1 and both are described as a first dose, although the "
  "second was given 21 days later; total = 2 counts only the Immunization entries. The import keeps the source "
  "value as OriginalDoseNumber, assigns DoseNumber by date order (1, then 2), adds the flag "
  "DUPLICATE_DOSE_NUMBER and sets ProcessingStatus to NeedsReview, so a verifying official decides. The Bundle "
  "also contains the NHS number and birth date, so OriginalContent is encrypted with Client-Side Field Level "
  "Encryption in production; the server stores only ciphertext.")

H("Scalability", 3)
P("VaccinationUpload, AuditLog and Notification are the fastest-growing data in PRS. They can be sharded "
  "horizontally (hashed PRSId for uploads and notifications, Timestamp for the audit log) without affecting the "
  "transactional core, which stays compact in SQL Server and can use read replicas for the public locator.")

H("3.4  Integrating the two models", 2)
P("The two stores form one design because they share identifiers, and those identifiers have the same names in "
  "every model:")
NAMES = [
    ("Person / PRSId", "Person.PRSId (PK)", "PRSId in VaccinationUpload, Notification; Target.PRSId in AuditLog"),
    ("UserAccount / UserId", "UserAccount.UserId (PK)", "UploadedByUserId; Actor.UserId"),
    ("Store / StoreId", "Store.StoreId (PK)", "Notification.StoreId"),
    ("GovernmentOfficial / OfficialId", "GovernmentOfficial.OfficialId (PK)", "Announcement.PublishedBy.OfficialId"),
    ("VaccinationUpload / UploadId", "VaccinationRecord.UploadId (CHAR(24))", "VaccinationUpload._id (ObjectId)"),
    ("VaccinationRecord / VaccinationRecordId", "VaccinationRecord.VaccinationRecordId (PK)", "ExtractedDoses[].VaccinationRecordId"),
]
T(["Concept (Figure 1)", "SQL Server (Figure 2)", "MongoDB (Figure 3)"], NAMES, widths=[5.0, 5.4, 6.2])
P("Table 4. Shared identifiers across the three models.", italic=True, size=10)
P("Because no single transaction spans both databases, writes follow a fixed order: the evidence is stored in "
  "MongoDB first, then the dose rows are written to SQL Server (idempotently), then the upload is updated with "
  "the new VaccinationRecordIds. If a step fails it can simply be repeated. A scheduled reconciliation job "
  "reports VaccinationRecords whose UploadId has no matching document and uploads left in Received state.")

# ================================================================ 4 CONCLUSION
H("4  Conclusion")
P("The PRS design separates a strict, transactional core from flexible, high-volume data. SQL Server holds "
  "identities, merchants, rules, purchases and normalised vaccination facts, where foreign keys, CHECK "
  "constraints and transactions guarantee that limits and schedules are applied consistently. MongoDB holds the "
  "data that does not fit tables well: FHIR reports whose structure varies, PDF and image certificates, the audit "
  "trail and notifications.")
P("Key design choices:")
B("Encryption is chosen per attribute: deterministic for identifiers that must be searched, randomised for the "
  "DOB, with only the birth-year digit kept in clear so tills can enforce schedules without decrypting anything.")
B("Imported data is never trusted blindly: the original evidence is kept unchanged, normalised doses are linked "
  "to it, and quality problems such as the duplicated dose number in the NHS sample are flagged for an official.")
B("Rules are data. Limits, schedules, vaccination policies, roles and permissions are rows that officials can "
  "change during a pandemic without changing the schema.")
B("Consistency across models is built in: one definition generates every diagram and the DDL, and the same "
  "names link the two stores.")
P("The main trade-off is that integrity between the two databases is maintained by the application and a "
  "reconciliation job rather than by foreign keys. This is accepted in exchange for storing evidence in its "
  "native form and scaling audit and upload data independently. Both models were implemented and tested, which "
  "shows the design supports the scenario's stories end to end.")

# ================================================================ REFERENCES
H("References")
for ref in [
    "NHS England (2026) Immunisation History - FHIR API. Available at: https://digital.nhs.uk/developer/api-catalogue/immunisation-history-fhir",
    "HL7 International (2023) FHIR R4 Immunization resource. Available at: https://hl7.org/fhir/R4/immunization.html",
    "Microsoft (2025) Always Encrypted. Available at: https://learn.microsoft.com/sql/relational-databases/security/encryption/always-encrypted-database-engine",
    "MongoDB Inc. (2025) Schema Validation; GridFS; Time Series Collections; Client-Side Field Level Encryption. Available at: https://www.mongodb.com/docs/manual/",
]:
    B(ref)

# ================================================================ APPENDIX
H("Appendix A  Implementation files and test evidence")
P("The models were implemented on Microsoft SQL Server 2022 and MongoDB 7.0 to check that they work. The files "
  "supplied with this report are:")
FILES = [
    ("prs_sqlserver_schema.sql", "Complete SQL Server DDL: 24 tables, constraints, indexes, seed data, procedures, view"),
    ("prs_mongodb_setup.js", "MongoDB collections, $jsonSchema validators, indexes and sample documents (mongosh)"),
    ("fhir_ingest.py", "Automatic import of a FHIR JSON report into both stores"),
    ("pdf_upload.py", "Stores a PDF/image certificate in GridFS and registers the upload"),
    ("test_scenarios.sql", "Purchase eligibility and constraint tests"),
    ("fig*.svg / fig*.png", "Full-resolution versions of every figure"),
]
T(["File", "Purpose"], FILES, widths=[5.0, 11.6])
P("Output of test_scenarios.sql (stories from section 2.3):")
CODE((BUILD / "test_sql_output.txt").read_text())
P("Output of fhir_ingest.py on the NHS sample, run twice:")
CODE((BUILD / "test_ingest_output.txt").read_text())

doc.save(OUTFILE)
print("wrote", OUTFILE)
